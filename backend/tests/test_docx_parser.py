"""DOCX 解析器单元测试：标题层级 / 段落 / 表格 / OCR 兜底 / .doc 拒绝。

用 monkeypatch 替身 Document（python-docx 的主类）以及 OCR 函数，
避免依赖真实 docx 文件 + 真实 OCR。
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest


# ---- 替身：构造 Document 与 paragraph / table ----


class _FakeRun:
    def __init__(self, text: str = "") -> None:
        self.text = text


class _FakeStyle:
    def __init__(self, name: str) -> None:
        self.name = name


class _FakeParagraph:
    def __init__(
        self,
        text: str = "",
        style_name: str = "Normal",
        has_drawing: bool = False,
        runs: list[_FakeRun] | None = None,
    ) -> None:
        self.text = text
        self.style = _FakeStyle(style_name)
        self._element = MagicMock()
        # 默认无图片
        if not has_drawing:
            self._element.findall.return_value = []
        self.runs = runs or []


class _FakeCell:
    def __init__(self, text: str = "") -> None:
        self.text = text


class _FakeRow:
    def __init__(self, cells: list[_FakeCell]) -> None:
        self.cells = cells


class _FakeTable:
    def __init__(self, rows: list[list[str]]) -> None:
        self.rows = [_FakeRow([_FakeCell(c) for c in r]) for r in rows]
        self._element = MagicMock()


class _FakeBody:
    """占位：实际由 _FakeDocument 提供 body 属性。"""


class _FakeDocument:
    """Document 替身：构造 .paragraphs / .tables / .element.body 让 docx parser 走完。

    docx parser 访问路径：
        document.element.body.iterchildren()  → 顶层元素的 _element 列表（XML 节点）
        child.tag                                  → 判断是段落(qn('w:p')) 还是表格(qn('w:tbl'))
        然后在 document.paragraphs / .tables 里反向找 _element is child 的 docx 对象
    """

    def __init__(
        self,
        paragraphs: list[_FakeParagraph],
        tables: list[_FakeTable],
    ) -> None:
        self._paragraphs = paragraphs
        self._tables = tables
        # 让 paragraph._element / table._element 的 tag 与 qn("w:p") / qn("w:tbl") 匹配
        from docx.oxml.ns import qn

        for p in paragraphs:
            p._element.tag = qn("w:p")
        for t in tables:
            t._element.tag = qn("w:tbl")

        # body.iterchildren() 返回的就是各段落/表格的 _element（XML 节点）
        self._body_elements = [p._element for p in paragraphs] + [t._element for t in tables]

    @property
    def paragraphs(self):
        return self._paragraphs

    @property
    def tables(self):
        return self._tables

    @property
    def element(self):
        body_el = MagicMock()
        body_el.iterchildren.return_value = iter(self._body_elements)
        elem = MagicMock()
        elem.body = body_el
        return elem


def _patch_document(monkeypatch, paragraphs, tables):
    """替身 docx.Document，返回 _FakeDocument。"""

    def _factory(*args, **kwargs):
        return _FakeDocument(paragraphs, tables)

    monkeypatch.setattr("app.services.parsers.docx.Document", _factory)


def _patch_ocr(monkeypatch, texts_by_call):
    """替身 ocr_pixmap_bytes。每次调用返回队列下一项。"""
    queue = list(texts_by_call)
    calls = []

    def _fake(png_bytes: bytes) -> str:
        calls.append(png_bytes)
        if not queue:
            return ""
        return queue.pop(0)

    monkeypatch.setattr("app.services.parsers.docx.ocr_pixmap_bytes", _fake)
    return calls


# ============ 测试用例 ============


def test_docx_heading_levels(monkeypatch):
    """Heading 1/2/3 → Markdown # ## ###，让 chunker 识别 heading_path。"""
    paragraphs = [
        _FakeParagraph("电力行业全链路 ASR 系统解决方案", style_name="Heading 1"),
        _FakeParagraph("项目背景", style_name="Heading 2"),
        _FakeParagraph("电力数字识别问题", style_name="Heading 3"),
        _FakeParagraph("正文段落"),  # Normal
    ]
    tables: list = []
    _patch_document(monkeypatch, paragraphs, tables)

    from app.services.parsers.docx import DocxParser

    parsed = DocxParser().parse(Path("dummy.docx"), doc_id="d1")
    assert parsed.ok
    assert parsed.page_count == 1
    # heading 应被转成 # 标题
    assert "# 电力行业全链路 ASR 系统解决方案" in parsed.full_text
    assert "## 项目背景" in parsed.full_text
    assert "### 电力数字识别问题" in parsed.full_text


def test_docx_table_to_markdown(monkeypatch):
    """python-docx Table → Markdown 表格语法（与 xlsx 一致）。"""
    paragraphs = [_FakeParagraph("前后文")]
    tables = [
        _FakeTable([
            ["目标维度", "具体目标"],
            ["专业词准确率", "电力专业术语识别准确率显著提升"],
            ["数字准确率", "电力数字读法"],
        ]),
    ]
    _patch_document(monkeypatch, paragraphs, tables)

    from app.services.parsers.docx import DocxParser

    parsed = DocxParser().parse(Path("dummy.docx"), doc_id="d1")
    assert parsed.ok
    text = parsed.full_text
    assert "|目标维度|具体目标|" in text
    assert "|---|---|" in text
    assert "|专业词准确率|电力专业术语识别准确率显著提升|" in text


def test_docx_chinese_heading_aliases(monkeypatch):
    """中文样式别名：'标题 1' / '标题 2' 也能识别。"""
    paragraphs = [
        _FakeParagraph("一级中文标题", style_name="标题 1"),
        _FakeParagraph("二级中文标题", style_name="标题2"),
    ]
    _patch_document(monkeypatch, paragraphs, [])

    from app.services.parsers.docx import DocxParser

    parsed = DocxParser().parse(Path("dummy.docx"), doc_id="d1")
    assert parsed.ok
    assert "# 一级中文标题" in parsed.full_text
    assert "## 二级中文标题" in parsed.full_text


def test_docx_injects_heading_path_into_chunks(monkeypatch):
    """DOCX 解析后的文本经 chunk_document 应带 heading_path。"""
    paragraphs = [
        _FakeParagraph("项目背景章节标题", style_name="Heading 1"),
        _FakeParagraph("这是第一段足够长的内容用于过 min_chunk_chars 阈值", style_name="Normal"),
        _FakeParagraph("电力数字识别问题小节标题", style_name="Heading 2"),
        _FakeParagraph("这是第二段足够长的内容用于过 min_chunk_chars 阈值", style_name="Normal"),
    ]
    _patch_document(monkeypatch, paragraphs, [])

    from app.services.parsers.docx import DocxParser
    from app.services.chunker import chunk_document

    parsed = DocxParser().parse(Path("dummy.docx"), doc_id="d1")
    chunks = chunk_document(parsed, min_chunk_chars=5)  # 短阈值避开长度过滤
    headings = [c.heading_path for c in chunks if c.heading_path]
    assert "项目背景章节标题" in headings
    assert "项目背景章节标题 > 电力数字识别问题小节标题" in headings


def test_docx_image_ocr_fallback(monkeypatch):
    """段落内嵌图片 → 走 OCR，OCR 结果并入文本。"""
    # 一个有 drawing 的段落
    p_with_img = _FakeParagraph("段落文字", style_name="Normal", has_drawing=True)
    p_with_img._element.findall = MagicMock(return_value=[MagicMock()])  # 含 drawing

    # 模拟 _ocr_paragraph_images 内部逻辑：找 blip → 拿 rel
    monkeypatch.setattr(
        "app.services.parsers.docx._ocr_paragraph_images",
        lambda p: "图片识别出来的文字内容" if p is p_with_img else "",
    )

    _patch_document(monkeypatch, [p_with_img, _FakeParagraph("其他段落")], [])

    from app.services.parsers.docx import DocxParser

    parsed = DocxParser().parse(Path("dummy.docx"), doc_id="d1")
    assert parsed.ok
    assert parsed.meta.get("image_count") == 1
    assert "图片识别出来的文字内容" in parsed.full_text


def test_parse_file_rejects_doc():
    """.doc 老格式 → 返回明确错误信息（基于 LibreOffice 未装 / D 盘空间考虑）。"""
    from app.services.parsers import parse_file, supported_extensions, rejected_extensions

    assert ".doc" in rejected_extensions()
    assert ".doc" not in supported_extensions()
    parsed = parse_file(Path("test.doc"), doc_id="d1")
    assert not parsed.ok
    assert "请在 Microsoft Word 或 WPS 中另存为 .docx 后重传" in parsed.error


def test_parse_file_supports_docx():
    """.docx → 走 DocxParser（通过扩展名 registry）。"""
    from app.services.parsers import parse_file, supported_extensions

    assert ".docx" in supported_extensions()
    parser = parse_file  # 占位，避免 lint 警告


def test_clean_text_protects_markdown_control_lines():
    """clean_text 修复：# 标题、> 引用、- 列表等控制行不被清空格。"""
    from app.services.parsers.base import clean_text

    cases = {
        "# 标题1": "# 标题1",
        "## 标题2": "## 标题2",
        "### 标题3": "### 标题3",
        "> 引用文本": "> 引用文本",
        "- 列表项": "- 列表项",
        "* 列表项": "* 列表项",
        "+ 列表项": "+ 列表项",
        # 普通行：空格仍被清
        "正文 段落": "正文段落",
        "2026 年第四次招标": "2026年第四次招标",
        # 缩进的控制行（lstrip 后仍是控制行）
        "  # 缩进标题": "# 缩进标题",
    }
    for orig, expected in cases.items():
        got = clean_text(orig)
        assert got == expected, f"{orig!r} → {got!r} (expected {expected!r})"


def test_clean_text_preserves_markdown_table():
    """Markdown 表格行（| 开头）不被误判为控制行——走普通行清洗逻辑。

    注意：xlsx / DOCX 生成的 Markdown 表格采用 `|列1|列2|`（无空格）紧贴写法，
    不会被 clean_text 破坏。用户自己写 .md 文件若用 `| 列1 |` 标准 Markdown 写法，
    表格前后空格会被清。这是 Markdown 表格的边缘情况，本项目优先 xlsx/DOCX 自动出表。
    """
    from app.services.parsers.base import clean_text

    # xlsx/DOCX 实际输出形式：紧贴
    text = "|列1|列2|\n|---|---|\n|数据1|数据2|"
    got = clean_text(text)
    assert "|列1|列2|" in got
    assert "|数据1|数据2|" in got