"""Excel 解析器（openpyxl → Markdown 表格）。

招标场景中 Excel 常用来发布中标候选人名单，本质是结构化表格，
转成 Markdown 表格后与 PDF 走同一套分块与检索逻辑。
每个工作表作为一个逻辑"页"，便于溯源。
"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook

from app.services.parsers.base import BaseParser, PageContent, ParsedDocument, clean_text

logger = logging.getLogger(__name__)


def _escape_cell(value) -> str:
    """单元格转字符串：去掉换行、转义竖线，避免破坏 Markdown 表格语法。"""
    if value is None:
        return ""
    text = str(value).strip()
    # 单元格内的换行会截断 Markdown 表格行
    text = " ".join(text.split())
    return text.replace("|", "\\|")


class XLSXParser(BaseParser):
    extensions = {".xlsx", ".xlsm"}

    def parse(self, path: Path, doc_id: str) -> ParsedDocument:
        doc = self._make(path, doc_id, parser="openpyxl")
        try:
            wb = load_workbook(str(path), data_only=True, read_only=True)
        except Exception as e:  # noqa: BLE001
            doc.error = f"{type(e).__name__}: {e}"
            logger.error("Excel 解析失败 %s: %s", path.name, e)
            return doc

        page_no = 1
        try:
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                md = self._to_markdown(rows)
                if not md:
                    continue
                text = clean_text(md)
                doc.pages.append(PageContent(page=page_no, text=text))
                page_no += 1
            doc.meta["sheets"] = [ws.title for ws in wb.worksheets]
        finally:
            wb.close()

        doc.meta["page_count"] = len(doc.pages)
        logger.info("Excel 解析完成 %s：%d 个表", path.name, len(doc.pages))
        return doc

    @staticmethod
    def _to_markdown(rows: list[tuple]) -> str:
        """二维数据转 Markdown 表格，跳过全空行。"""
        data = [[_escape_cell(c) for c in row] for row in rows]
        data = [r for r in data if any(c for c in r)]
        if not data:
            return ""

        # 统一列数，避免 Markdown 表格错位
        width = max(len(r) for r in data)
        for r in data:
            r.extend([""] * (width - len(r)))

        header = data[0]
        lines = ["|" + "|".join(header) + "|", "|" + "|".join(["---"] * width) + "|"]
        for r in data[1:]:
            lines.append("|" + "|".join(r) + "|")
        return "\n".join(lines)
